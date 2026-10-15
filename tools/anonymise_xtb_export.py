"""Anonymise an XTB xStation account-history export (xlsx) for use as a test fixture.

- The account number -> 12345678 wherever it appears. In XTB's export it is in B1 of every
  sheet and also fills the cash "ID" and "Position ID" columns and the deposit comments.
- Open-lot position IDs remapped to new 10-digit numbers, consistently.
- Whole-share fills (e.g. "OPEN BUY 1/1.2345") cannot be scaled while staying whole, so all
  volumes are written with at least 4 decimals; otherwise those rows would keep real amounts.
- Every quantity and money value is scaled by one random factor k (never printed or stored).
  Amounts are recomputed from the rounded quantity so that amount / (quantity x price), i.e.
  XTB's EUR->PLN conversion, stays as in the original. Prices, dates and percentages are kept.
- Totals and summary values are recomputed from the anonymised rows.

Usage (from the repo root; needs uv):

    uv run --no-project --with openpyxl python -I tools/anonymise_xtb_export.py IN.xlsx OUT.xlsx

Check the output by eye before committing it: this only handles the layout seen so far.
"""

import random
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

ACCOUNT = 12345678

K = Decimal(str(round(random.SystemRandom().uniform(0.3, 0.7), 6)))


def dec(v):
    return Decimal(repr(v)) if isinstance(v, float) else Decimal(v)


def places(text: str) -> int:
    return len(text.split(".")[1]) if "." in text else 0


def q(x: Decimal, dp: int) -> Decimal:
    return x.quantize(Decimal(1).scaleb(-dp), rounding=ROUND_HALF_UP)


def scale_volume(v: Decimal, dp: int) -> Decimal:
    """Scaled volume with at least 4 decimals (so whole shares change too), never zero."""
    dp = max(dp, 4)
    return max(q(v * K, dp), Decimal(1).scaleb(-dp))


def num(x: Decimal):
    """Back to a cell value: int when whole, else float (written as its shortest repr)."""
    return int(x) if x == x.to_integral_value() and "." not in str(x) else float(x)


class IdMap:
    """Original id -> new id with the same number of digits, sequential per map."""

    def __init__(self, first_digit: str):
        self.first_digit = first_digit
        self.ids: dict[str, str] = {}

    def __call__(self, original):
        key = str(original)
        if key not in self.ids:
            n = len(self.ids) + 1
            self.ids[key] = self.first_digit + str(n).zfill(len(key) - 1)
        return int(self.ids[key]) if isinstance(original, int) else self.ids[key]


lot_ids = IdMap("3")

BUY = re.compile(r"^(OPEN BUY (?:\S+ )?)(?:([\d.]+)/)?([\d.]+)( @ [\d.]+)$")


def anonymise_cash(ws):
    rows = [r for r in ws.iter_rows(min_row=6) if r[0].value]
    for r in rows:
        kind = r[0].value
        if kind == "Stock purchase":
            m = BUY.match(r[7].value)
            if not m:
                raise SystemExit(f"unexpected purchase comment format in row {r[0].row}")
            prefix, filled_s, second_s, suffix = m.groups()
            if filled_s:  # "filled/ordered"
                filled, ordered = dec(filled_s), dec(second_s)
                new_filled = scale_volume(filled, places(filled_s))
                new_ordered = max(scale_volume(ordered, places(second_s)), new_filled)
                volume_text = f"{new_filled}/{new_ordered}"
            else:
                filled = dec(second_s)
                new_filled = scale_volume(filled, places(second_s))
                volume_text = f"{new_filled}"
            r[7].value = f"{prefix}{volume_text}{suffix}"
            amount = dec(r[5].value)
            r[5].value = num(q(amount * new_filled / filled, 2))
        elif kind == "IKZE deposit":
            amount = dec(r[5].value)
            dp = 0 if amount == amount.to_integral_value() else 2
            r[5].value = num(q(amount * K, dp))
        elif kind == "Total":
            continue
        else:
            raise SystemExit(f"unknown cash operation type {kind!r} in row {r[0].row}")

    data = [r for r in rows if r[0].value != "Total"]
    total_row = next(r for r in rows if r[0].value == "Total")
    original_total_non_negative = dec(total_row[5].value) >= 0
    total = sum(dec(r[5].value) for r in data)
    if original_total_non_negative and total < 0:
        # Rounding pushed the cash balance below zero: top up the last deposit.
        last_deposit = [r for r in data if r[0].value == "IKZE deposit"][-1]
        last_deposit[5].value = num(dec(last_deposit[5].value) - total.to_integral_value("ROUND_FLOOR"))
        total = sum(dec(r[5].value) for r in data)
    total_row[5].value = num(q(total, 2))


def anonymise_open(ws):
    lots = [r for r in ws.iter_rows(min_row=10) if r[0].value == "IKZE" and str(r[1].value).isdigit()]
    for r in lots:
        r[1].value = lot_ids(r[1].value)
        vol_text = repr(r[5].value) if isinstance(r[5].value, float) else str(r[5].value)
        volume = dec(r[5].value)
        new_volume = scale_volume(volume, places(vol_text))
        ratio = new_volume / volume
        r[5].value = num(new_volume)
        for i in (6, 13, 14):  # value, net profit, gross profit
            if r[i].value is not None:
                r[i].value = num(q(dec(r[i].value) * ratio, 2))
    summary = next(r for r in ws.iter_rows(min_row=10) if r[1].value == "FTSE All-World")
    summary[5].value = num(sum(dec(r[5].value) for r in lots))
    for i in (6, 13, 14):
        summary[i].value = num(q(sum(dec(r[i].value) for r in lots), 2))
    open_dp = places(repr(summary[8].value))
    weighted = sum(dec(r[5].value) * dec(r[8].value) for r in lots) / dec(summary[5].value)
    summary[8].value = num(q(weighted, open_dp))
    ws["C5"].value = summary[6].value
    ws["C6"].value = summary[13].value


def replace_account(wb, account: int):
    """Every cell equal to the account number, and every occurrence inside text."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == account or cell.value == str(account):
                    cell.value = ACCOUNT if isinstance(cell.value, int) else str(ACCOUNT)
                elif isinstance(cell.value, str) and str(account) in cell.value:
                    cell.value = cell.value.replace(str(account), str(ACCOUNT))


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    account = wb.worksheets[0]["B1"].value
    if not isinstance(account, int):
        raise SystemExit("B1 of the first sheet is not the account number")
    anonymise_cash(wb["Cash Operations"])
    anonymise_open(wb["Open Positions"])
    replace_account(wb, account)
    # Workbook metadata could carry identifying text; reset it.
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    wb.properties.title = None
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
